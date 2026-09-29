import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import docx

def create_sample_files():
    sample_dir = os.path.join(os.path.dirname(__file__), "sample_data")
    os.makedirs(sample_dir, exist_ok=True)

    # 1. GENERATE CATALOG EXCEL: sample_data/danh_muc_dinh_muc_mau.xlsx
    wb_catalog = openpyxl.Workbook()
    ws_cat = wb_catalog.active
    ws_cat.title = "DinhMuc"

    # Explanatory text above table (should be filtered out by smart extraction)
    ws_cat["A1"] = "BỘ XÂY DỰNG - ĐỊNH MỨC DỰ TOÁN XÂY DỰNG CÔNG TRÌNH"
    ws_cat["A2"] = "(Ban hành kèm theo Thông tư số 12/2021/TT-BXD của Bộ Xây dựng)"
    ws_cat["A3"] = "Phần thuyết minh: Định mức dự toán xây dựng công trình quy định mức hao phí cần thiết về vật liệu..."

    # Table Header at row 5
    headers = ["STT", "Mã hiệu ĐM", "Tên công tác xây lắp", "Đơn vị tính", "Ghi chú"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws_cat.cell(row=5, column=col_idx, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")

    catalog_rows = [
        (1, "AB.123", "Đào móng công trình bằng máy đào 0.8m3, đất cấp II", "100m3", "Mã ngắn 3 số -> đệm thành AB.12300"),
        (2, "AF.1234", "Bê tông lót móng đá 4x6 mác 100", "m3", "Mã ngắn 4 số -> đệm thành AF.12340"),
        (3, "AF.21111", "Bê tông móng đổ bằng thủ công đá 1x2 mác 200", "m3", "Mã chuẩn 5 số"),
        (4, "AF.22222", "Bê tông cột vách đá 1x2 mác 250", "m3", "Mã chuẩn 5 số"),
        (5, "AF.33333", "Bê tông dầm sàn đá 1x2 mác 250", "m3", "Mã chuẩn 5 số"),
        (6, "AF.61111", "Ván khuôn móng thép", "100m2", "Mã chuẩn 5 số"),
        (7, "AF.62222", "Ván khuôn cột dầm sàn thép", "100m2", "Mã chuẩn 5 số"),
        (8, "AF.63333", "Cốt thép dầm sàn đường kính <= 10mm", "tấn", "Mã chuẩn 5 số"),
        (9, "AF.64444", "Cốt thép dầm sàn đường kính <= 18mm", "tấn", "Mã chuẩn 5 số"),
        (10, "AK.11111", "Xây tường gạch ống 8x8x18 vữa xi măng mác 75", "m3", "Mã chuẩn 5 số"),
        (11, "AK.22222", "Trát tường trong vữa xi măng mác 75 dày 1.5cm", "m2", "Mã chuẩn 5 số"),
        (12, "AK.33333", "Lát gạch granite nhân tạo 600x600", "m2", "Mã chuẩn 5 số"),
        (13, "BA.111", "Lắp đặt dây dẫn điện đơn ruột đồng", "100m", "Mã ngắn 3 số -> BA.11100"),
        (14, "BA.2222", "Lắp đặt tủ điện phân phối chiếu sáng", "cái", "Mã ngắn 4 số -> BA.22220"),
        (15, "BA.33333", "Lắp đặt đèn LED panel 600x600 48W", "bộ", "Mã chuẩn 5 số"),
    ]

    for row_idx, r_data in enumerate(catalog_rows, 6):
        for col_idx, val in enumerate(r_data, 1):
            ws_cat.cell(row=row_idx, column=col_idx, value=val)

    # Explanatory notes at the bottom
    last_r = len(catalog_rows) + 7
    ws_cat.cell(row=last_r, column=1, value="Ghi chú thêm: Bảng định mức trên có giá trị áp dụng từ năm 2024.")

    catalog_excel_path = os.path.join(sample_dir, "danh_muc_dinh_muc_mau.xlsx")
    wb_catalog.save(catalog_excel_path)
    print(f"Created: {catalog_excel_path}")

    # 2. GENERATE CATALOG WORD (.docx): sample_data/danh_muc_dinh_muc_mau.docx
    doc = docx.Document()
    doc.add_heading("BẢNG TỔNG HỢP DANH MỤC CÔNG VIỆC CHUẨN", level=1)
    doc.add_paragraph("Tài liệu nội bộ dự án: Danh mục mã hiệu quy chuẩn phục vụ quản lý dự án.")
    doc.add_paragraph("Lưu ý: Các mã công việc phải tuân thủ nghiêm ngặt định dạng quy định của Corebox.")

    doc_table = doc.add_table(rows=1, cols=4)
    hdr_cells = doc_table.rows[0].cells
    hdr_cells[0].text = "STT"
    hdr_cells[1].text = "Mã CV"
    hdr_cells[2].text = "Tên công việc"
    hdr_cells[3].text = "Đơn vị"

    for r in catalog_rows[:8]:
        row_cells = doc_table.add_row().cells
        row_cells[0].text = str(r[0])
        row_cells[1].text = r[1]
        row_cells[2].text = r[2]
        row_cells[3].text = r[3]

    doc.add_paragraph("Ký tên xác nhận: Trưởng phòng Kỹ thuật BQLDA")
    catalog_docx_path = os.path.join(sample_dir, "danh_muc_dinh_muc_mau.docx")
    doc.save(catalog_docx_path)
    print(f"Created: {catalog_docx_path}")

    # 3. GENERATE INSPECTION EXCEL: sample_data/du_toan_kiem_tra_DGTH.xlsx
    # MUST contain sheet 'ĐGTH'
    # Multiple sections where STT restarts from 1
    wb_inspect = openpyxl.Workbook()
    # Remove default sheet and create exact 'ĐGTH'
    wb_inspect.remove(wb_inspect.active)
    ws_dgth = wb_inspect.create_sheet(title="ĐGTH")

    ws_dgth["A1"] = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    ws_dgth["A2"] = "DỰ TOÁN CÔNG TRÌNH: TRỤ SỞ LÀM VIỆC BQLDA"
    ws_dgth["A3"] = "BẢNG ĐƠN GIÁ TỔNG HỢP (ĐGTH)"

    current_r = 5

    # SECTION 1: MÓNG CÔNG TRÌNH
    ws_dgth.cell(row=current_r, column=1, value="Hạng mục 1: PHẦN MÓNG CÔNG TRÌNH").font = Font(bold=True, size=12, color="1E3A8A")
    current_r += 1

    headers_sec = ["STT", "Mã hiệu", "Nội dung công việc", "Đơn vị tính", "Khối lượng", "Đơn giá", "Thành tiền"]
    for c_i, h in enumerate(headers_sec, 1):
        cell = ws_dgth.cell(row=current_r, column=c_i, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    current_r += 1

    # Items for Section 1:
    sec1_items = [
        (1, "AB.123", "Đào móng công trình bằng máy đào 0.8m3, đất cấp II", "100m3", 15.5, 1200000, 18600000),  # Valid (matches padded AB.12300)
        (2, "AF.1234", "Bê tông lót móng đá 4x6 mác 100", "m3", 45.2, 850000, 38420000),                      # Valid (matches padded AF.12340)
        (3, "AF.21111", "Bê tông móng đổ bằng thủ công đá 1x2 mác 200", "m3", 110.0, 1250000, 137500000),     # Valid
        (4, "AF.99999", "Mã công việc không có trong kho", "m3", 10.0, 900000, 9000000),                       # ERROR: Not found!
        (5, "AF.61111", "Ván khuôn móng thép", "100m2", 25.0, 3200000, 80000000),                              # Valid
    ]
    for it in sec1_items:
        for c_i, val in enumerate(it, 1):
            ws_dgth.cell(row=current_r, column=c_i, value=val)
        current_r += 1

    current_r += 2

    # SECTION 2: THÂN VÀ HOÀN THIỆN (STT restarts from 1)
    ws_dgth.cell(row=current_r, column=1, value="Hạng mục 2: PHẦN THÂN VÀ HOÀN THIỆN").font = Font(bold=True, size=12, color="1E3A8A")
    current_r += 1

    for c_i, h in enumerate(headers_sec, 1):
        cell = ws_dgth.cell(row=current_r, column=c_i, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    current_r += 1

    sec2_items = [
        (1, "AF.22222", "Bê tông cột vách đá 1x2 mác 250", "m3", 85.0, 1350000, 114750000),                   # Valid
        (2, "AF.33333", "Bê tông dầm sàn đá 1x2 mác 250", "m3", 140.0, 1350000, 189000000),                   # Valid
        (3, "AK.999", "Mã lỗi ngắn không có trong kho", "m2", 50.0, 150000, 7500000),                          # ERROR: Not found!
        (4, "AK.11111", "Xây tường gạch ống 8x8x18 vữa xi măng mác 75", "m3", 65.0, 1100000, 71500000),      # Valid
        (5, "AK.33333", "Lát gạch granite nhân tạo 600x600", "m2", 320.0, 280000, 89600000),                   # Valid
    ]
    for it in sec2_items:
        for c_i, val in enumerate(it, 1):
            ws_dgth.cell(row=current_r, column=c_i, value=val)
        current_r += 1

    current_r += 2

    # SECTION 3: ĐIỆN VÀ CHIẾU SÁNG (STT restarts from 1)
    ws_dgth.cell(row=current_r, column=1, value="Hạng mục 3: HỆ THỐNG ĐIỆN VÀ CHIẾU SÁNG").font = Font(bold=True, size=12, color="1E3A8A")
    current_r += 1

    for c_i, h in enumerate(headers_sec, 1):
        cell = ws_dgth.cell(row=current_r, column=c_i, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    current_r += 1

    sec3_items = [
        (1, "BA.111", "Lắp đặt dây dẫn điện đơn ruột đồng", "100m", 35.0, 650000, 22750000),                  # Valid (matches padded BA.11100)
        (2, "BA.2222", "Lắp đặt tủ điện phân phối chiếu sáng", "cái", 4, 3500000, 14000000),                   # Valid (matches padded BA.22220)
        (3, "BA.33333", "Lắp đặt đèn LED panel 600x600 48W", "bộ", 60, 420000, 25200000),                     # Valid
        (4, "XX.88888", "Thiết bị điện sai mã", "cái", 10, 500000, 5000000),                                   # ERROR: Not found!
    ]
    for it in sec3_items:
        for c_i, val in enumerate(it, 1):
            ws_dgth.cell(row=current_r, column=c_i, value=val)
        current_r += 1

    inspect_excel_path = os.path.join(sample_dir, "du_toan_kiem_tra_DGTH.xlsx")
    wb_inspect.save(inspect_excel_path)
    print(f"Created: {inspect_excel_path}")

if __name__ == "__main__":
    create_sample_files()
